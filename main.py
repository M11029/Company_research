# import os
# import streamlit as st
# import requests
# from dotenv import load_dotenv
# import base64
# import json
# from docx import Document
# from io import BytesIO
# from openai import OpenAI
# from streamlit_option_menu import option_menu  # Import the option_menu library

# # Load environment variables from a .env file for development purposes
# load_dotenv()

# # Set API keys using environment variables
# companies_house_api_key = os.getenv("COMPANIES_HOUSE_API_KEY")
# openai_api_key = os.getenv("OPENAI_API_KEY")
# news_api_key = os.getenv("NEWS_API_KEY")

# # Initialize OpenAI API key
# client = OpenAI(api_key=openai_api_key)

# # Set page configuration
# st.set_page_config(page_title="Company Info App", layout="wide")

# if 'current_tab' not in st.session_state:
#     st.session_state['current_tab'] = 'Welcome'


# def main():
#     """
#     Main function to run the Streamlit app.
#     Initializes session state and handles navigation between tabs.
#     """
#     # Initialize session state for the current tab if not already set
#     if 'current_tab' not in st.session_state:
#         st.session_state['current_tab'] = 'Welcome'

#     # List of tabs in the app
#     tabs = ['Welcome', 'Enter Company Name', 'Results', 'Other Project']
#     icons = ['house', 'building', 'bar-chart', 'info-circle']

#     # Navigation pane in the sidebar using option_menu
#     with st.sidebar:
#         st.title("Navigation")
#         selected_tab = option_menu(
#             menu_title=None,  # Hide the menu title
#             options=tabs,
#             icons=icons,
#             menu_icon="cast",
#             default_index=tabs.index(st.session_state['current_tab']),
#             styles={
#                 "container": {"padding": "0!important", "background-color": "#f0f0f0"},
#                 "icon": {"color": "black", "font-size": "18px"},
#                 "nav-link": {
#                     "font-size": "16px",
#                     "text-align": "center",
#                     "margin": "5px",
#                     "color": "black",
#                     "background-color": "#f0f0f0",
#                     "border-radius": "5px",
#                 },
#                 "nav-link-hover": {"background-color": "#e0e0e0"},
#                 "nav-link-selected": {"background-color": "#007bff", "color": "white"},
#             },
#             key='navigation'
#         )
#         st.session_state['current_tab'] = selected_tab



#     # Display content based on current tab
#     if st.session_state['current_tab'] == 'Welcome':
#         welcome_tab()
#     elif st.session_state['current_tab'] == 'Enter Company Name':
#         enter_company_name_tab()
#     elif st.session_state['current_tab'] == 'Results':
#         results_tab()
#     elif st.session_state['current_tab'] == 'Other Project':
#         other_project_tab()

# def navigate_tabs(direction):
#     """
#     Function to navigate between tabs using 'Next' and 'Previous' buttons.
#     Updates the current tab in session state.
#     """
#     tabs = ['Welcome', 'Enter Company Name', 'Results', 'Other Project']
#     current_index = tabs.index(st.session_state['current_tab'])
#     if direction == 'next': 
#     #and current_index < len(tabs) - 1:
#         st.session_state['current_tab'] = tabs[current_index + 1]
#     elif direction == 'previous': 
#     #and current_index > 0:
#         st.session_state['current_tab'] = tabs[current_index - 1]
#     #st.experimental_rerun()
#     st.rerun()


# def welcome_tab():
#     """
#     Displays the content for the 'Welcome' tab.
#     Introduces the app and its features.
#     """
#     st.title("Welcome to the Company Info App")
#     st.write("""
#     ### Discover Detailed Company Information Easily!

#     Welcome to the **Company Info App**! This application allows you to effortlessly search for companies registered in the UK using the Companies House API.

#     **What Can You Do With This App?**

#     - **Search Companies:** Enter the names of companies you're interested in and select the correct one from a list of suggestions.
#     - **View Detailed Information:** Access comprehensive details including company status, incorporation date, registered office address, officers, and more.
#     - **Stay Updated:** Read the latest news articles related to the company to stay informed about recent developments.
#     - **Generate Summaries:** Get concise summaries of the company's information and recent news.
#     - **Download Reports:** Export all the gathered information into a nicely formatted Word document for your records or presentations.

#     **Get Started Now!**

#     Navigate to the **Enter Company Name** tab to begin your search and explore detailed company profiles.
#     """)

#     # Navigation buttons
#     col1, col2, col3 = st.columns([1, 1, 1])
#     with col3:
#         if st.button("Next", key="welcome_next"):
#             navigate_tabs('next')

# def enter_company_name_tab():
#     """
#     Displays the content for the 'Enter Company Name' tab.
#     Allows users to input company names and initiates the selection process.
#     """
#     st.title("Enter Company Name")

#     # Input for company names
#     company_input = st.text_area("Enter company names (one per line):", key='company_input')

#     if st.button("Search Companies", key="search_companies"):
#         company_names = company_input.strip().split('\n')
#         company_names = [name.strip() for name in company_names if name.strip()]

#         if company_names:
#             # Initialize session state variables for the selection process
#             st.session_state['company_names'] = company_names
#             st.session_state['confirmed_companies'] = {}
#             st.session_state['current_company_index'] = 0
#             st.session_state['company_selections'] = {}
#             st.session_state['selection_complete'] = False
#             st.session_state['company_selection_started'] = True
#             # Rerun the app to update the interface
#             st.rerun()
#             #st.experimental_rerun()
#         else:
#             st.warning("Please enter at least one company name.")

#     # If the selection process has started, continue it
#     if st.session_state.get('company_selection_started', False):
#         select_companies()

#     # Navigation buttons
#     col1, col2, col3 = st.columns([1, 1, 1])
#     with col1:
#         if st.button("Previous", key="enter_previous"):
#             navigate_tabs('previous')
#     with col3:
#         if st.button("Next", key="enter_next"):
#             navigate_tabs('next')

# def select_companies():
#     """
#     Handles the company selection process.
#     Displays search results and allows the user to confirm the correct company.
#     """
#     st.title("Select Companies")

#     company_names = st.session_state['company_names']
#     total_companies = len(company_names)
#     current_index = st.session_state['current_company_index']

#     # Progress bar indicating how much of the selection process is complete
#     progress = int((current_index / total_companies) * 100)
#     st.progress(progress)

#     if current_index < total_companies:
#         company_name = company_names[current_index]
#         st.subheader(f"Results for '{company_name}'")

#         # Search Companies House API
#         if not companies_house_api_key:
#             st.error("Companies House API key is missing.")
#             return

#         # Prepare authorization for Companies House API
#         auth_string = f"{companies_house_api_key}:"
#         headers = {
#             'Authorization': 'Basic ' + base64.b64encode(auth_string.encode()).decode()
#         }
#         params = {'q': company_name}

#         with st.spinner('Fetching company data...'):
#             try:
#                 response = requests.get(
#                     'https://api.company-information.service.gov.uk/search/companies',
#                     headers=headers,
#                     params=params
#                 )
#                 response.raise_for_status()
#             except requests.exceptions.RequestException as e:
#                 st.error(f"Error fetching data for '{company_name}': {e}")
#                 # Move to next company
#                 st.session_state['company_selections'][company_name] = None
#                 st.session_state['current_company_index'] += 1
#                 st.rerun()
#                 #st.experimental_rerun()
#                 return

#         data = response.json()
#         items = data.get('items', [])

#         if not items:
#             st.write(f"No results found for '{company_name}'.")
#             # Move to next company
#             st.session_state['company_selections'][company_name] = None
#             st.session_state['current_company_index'] += 1
#             st.rerun()
#             #st.experimental_rerun()
#             return

#         # Show top 5 results as a dropdown
#         options = []
#         for item in items[:5]:
#             options.append(f"{item.get('title')} (Company Number: {item.get('company_number')})")

#         # Use a form to group the selectbox and the confirm button
#         with st.form(key=f"form_{company_name}"):
#             selected = st.selectbox(
#                 f"Select the correct company for '{company_name}':", options, key=f"select_{company_name}"
#             )
#             confirm = st.form_submit_button("Confirm Selection")

#         if confirm:
#             # Store the selected company number
#             selected_index = options.index(selected)
#             selected_company = items[selected_index]
#             company_number = selected_company.get('company_number')

#             # Save the selection in session state
#             st.session_state['company_selections'][company_name] = company_number
#             # Update current company index
#             st.session_state['current_company_index'] += 1
#             # Check if selection is complete
#             if st.session_state['current_company_index'] >= total_companies:
#                 st.session_state['confirmed_companies'] = st.session_state['company_selections']
#                 st.session_state['selection_complete'] = True
#                 st.session_state['company_selection_started'] = False
#                 st.success("Company selection complete!")
#                 # Navigate to Results tab
#             st.session_state['current_tab'] = 'Results'
#             # Rerun the app to update
#             st.rerun()
#             #st.experimental_rerun()
#     else:
#         # All companies have been processed
#         st.success("Company selection complete!")
#         st.session_state['confirmed_companies'] = st.session_state['company_selections']
#         st.session_state['selection_complete'] = True
#         st.session_state['company_selection_started'] = False
#         # Navigate to Results tab
#         st.session_state['current_tab'] = 'Results'
#         st.rerun()
#         #st.experimental_rerun()

# def results_tab():
#     """
#     Displays the results for the selected companies.
#     Shows detailed information and provides options to download reports.
#     """
#     st.title("Results")

#     if 'confirmed_companies' not in st.session_state or not st.session_state['confirmed_companies']:
#         st.warning("No confirmed companies found. Please go to the 'Enter Company Name' tab to search for companies.")
#         return

#     companies = st.session_state['confirmed_companies']

#     for company_name, company_number in companies.items():
#         if company_number is None:
#             st.header(f"{company_name} - No Data Available")
#             continue
#         st.header(f"{company_name} (Company Number: {company_number})")
#         display_company_details(company_name, company_number)

#     # Navigation buttons
#     col1, col2, col3 = st.columns([1, 1, 1])
#     with col1:
#         if st.button("Previous", key="results_previous"):
#             navigate_tabs('previous')
#     with col3:
#         if st.button("Next", key="results_next"):
#             navigate_tabs('next')

# def display_company_details(company_name, company_number):
#     """
#     Fetches and displays detailed information for a single company.
#     Includes important information, additional data, recent news, and a summary.
#     """
#     # Fetch company details from Companies House API
#     if not companies_house_api_key:
#         st.error("Companies House API key is missing.")
#         return

#     # Prepare authorization for Companies House API
#     auth_string = f"{companies_house_api_key}:"
#     headers = {
#         'Authorization': 'Basic ' + base64.b64encode(auth_string.encode()).decode()
#     }

#     with st.spinner('Fetching company details...'):
#         try:
#             response = requests.get(
#                 f'https://api.company-information.service.gov.uk/company/{company_number}',
#                 headers=headers
#             )
#             response.raise_for_status()
#         except requests.exceptions.RequestException as e:
#             st.error(f"Error fetching details for '{company_name}': {e}")
#             return

#     company_data = response.json()

#     # Display important information in a styled format
#     st.subheader("Important Information")
#     important_info = {
#         'Company Status': company_data.get('company_status'),
#         'Incorporation Date': company_data.get('date_of_creation'),
#         'Registered Office Address': company_data.get('registered_office_address')
#     }

#     for key, value in important_info.items():
#         st.write(f"**{key}:** {value}")

#     # Additional Information
#     st.subheader("Additional Information")
#     additional_info_options = ['SIC Codes', 'Accounts', 'Officers']
#     selected_info = st.multiselect(
#         "Select additional information to view:", additional_info_options, key=f"info_{company_name}"
#     )

#     if 'SIC Codes' in selected_info:
#         sic_codes = company_data.get('sic_codes', [])
#         st.write("**SIC Codes:**")
#         for code in sic_codes:
#             st.write(f"- {code}")
#     else:
#         sic_codes = []

#     if 'Accounts' in selected_info:
#         accounts = company_data.get('accounts', {})
#         st.write("**Accounts:**")
#         st.json(accounts)
#     else:
#         accounts = {}

#     if 'Officers' in selected_info:
#         # Fetch officers from Companies House API
#         with st.spinner('Fetching officer information...'):
#             try:
#                 officers_response = requests.get(
#                     f'https://api.company-information.service.gov.uk/company/{company_number}/officers',
#                     headers=headers
#                 )
#                 officers_response.raise_for_status()
#             except requests.exceptions.RequestException as e:
#                 st.write("Could not retrieve officers information.")
#                 officers = []
#             else:
#                 officers_data = officers_response.json()
#                 officers = officers_data.get('items', [])
#                 st.write("**Officers:**")
#                 for officer in officers:
#                     st.write(f"- {officer.get('name')} ({officer.get('officer_role')})")
#     else:
#         officers = []

#     # Recent news section using News API
#     st.subheader("Recent News")
#     news_params = {
#         'q': company_name,
#         'apiKey': news_api_key,
#         'pageSize': 5,
#         'sortBy': 'publishedAt'
#     }

#     with st.spinner('Fetching recent news...'):
#         news_response = requests.get('https://newsapi.org/v2/everything', params=news_params)
#         if news_response.status_code == 200:
#             news_data = news_response.json()
#             articles = news_data.get('articles', [])
#             if articles:
#                 for article in articles:
#                     st.write(f"- [{article.get('title')}]({article.get('url')})")
#             else:
#                 st.write("No recent news found.")
#         else:
#             st.write("Error fetching news articles.")
#             articles = []
#     # Generate 250-word summary using OpenAI API
#     st.subheader("Summary")
#     with st.spinner('Generating summary...'):
#         summary = generate_summary(company_data, articles)
#         st.write(summary)

#     # Download button for Word document
#     if st.button("Download Report as Word Document", key=f"download_{company_name}"):
#         generate_word_report(
#             company_name=company_name,
#             important_info=important_info,
#             selected_info=selected_info,
#             sic_codes=sic_codes,
#             accounts=accounts,
#             officers=officers,
#             articles=articles,
#             summary=summary  # Use the generated summary
#         )

# def generate_summary(company_data, articles):
#     """
#     Generates a summary of the company information and recent news using OpenAI's GPT-4.
#     """
#     # Prepare the content for the summary
#     content = f"Company Information:\n"
#     content += f"Status: {company_data.get('company_status')}\n"
#     content += f"Incorporation Date: {company_data.get('date_of_creation')}\n"
#     content += f"Registered Office Address: {company_data.get('registered_office_address')}\n\n"

#     content += "Recent News Articles:\n"
#     for article in articles:
#         content += f"- {article.get('title')}: {article.get('description')}\n"

#     # Use OpenAI's GPT-4 to generate a 250-word summary
#     try:
#         response = client.chat.completions.create(
#             model="gpt-4",
#             messages=[
#                 {"role": "system", "content": "You are a helpful assistant that summarizes company information."},
#                 {"role": "user", "content": f"Provide a 250-word summary for the following company information and news articles:\n{content}",}
#             ],
#             max_tokens=400,
#             temperature=0.7
#         )
#         summary = response.choices[0].message.content.strip()

#     except Exception as e:
#         summary = "Could not generate summary."
#         st.error(f"Error generating summary: {e}")

#     return summary

# def generate_word_report(company_name, important_info, selected_info, sic_codes, accounts, officers, articles, summary):
#     """
#     Generates a Word document report for the company and provides a download link.
#     """
#     doc = Document()
#     doc.add_heading(f"{company_name} Report", 0)

#     # Add important information
#     doc.add_heading("Important Information", level=1)
#     for key, value in important_info.items():
#         doc.add_paragraph(f"{key}: {value}")

#     # Add additional information
#     if selected_info:
#         doc.add_heading("Additional Information", level=1)
#         if 'SIC Codes' in selected_info and sic_codes:
#             doc.add_heading("SIC Codes", level=2)
#             for code in sic_codes:
#                 doc.add_paragraph(code, style='List Bullet')

#         if 'Accounts' in selected_info and accounts:
#             doc.add_heading("Accounts", level=2)
#             doc.add_paragraph(json.dumps(accounts, indent=2))

#         if 'Officers' in selected_info and officers:
#             doc.add_heading("Officers", level=2)
#             for officer in officers:
#                 doc.add_paragraph(
#                     f"{officer.get('name')} ({officer.get('officer_role')})", style='List Bullet'
#                 )

#     # Add recent news
#     doc.add_heading("Recent News", level=1)
#     if articles:
#         for article in articles:
#             p = doc.add_paragraph(style='List Bullet')
#             p.add_run(article.get('title')).bold = True
#             p.add_run(f"\n{article.get('description')}\n")
#             p.add_run(f"Link: {article.get('url')}")
#     else:
#         doc.add_paragraph("No recent news articles available.")

#     # Add summary
#     doc.add_heading("Summary", level=1)
#     doc.add_paragraph(summary)

#     # Save the document to a BytesIO object
#     buffer = BytesIO()
#     doc.save(buffer)
#     buffer.seek(0)

#     # Provide download link
#     st.download_button(
#         label="Download Word Document",
#         data=buffer,
#         file_name=f"{company_name}_Report.docx",
#         mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
#     )

# def other_project_tab():
#     """
#     Displays the content for the 'Other Project' tab.
#     Provides a thank-you message and a link to another project.
#     """
#     st.title("Other Project")
#     st.write("Thank you for trying out my app. I hope you enjoyed it.")
#     st.write("Please try out my other project:")
#     st.markdown("[Click here to visit my other project](https://app2py-ta4gt3rhhab4to4cqz9hgi.streamlit.app/)")

#     # Navigation buttons
#     col1, col2, col3 = st.columns([1, 1, 1])
#     with col1:
#         if st.button("Previous", key="other_previous"):
#             navigate_tabs('previous')

# if __name__ == "__main__":
#     main()
import os
import streamlit as st
import requests
from dotenv import load_dotenv
import base64
import json
from docx import Document
from io import BytesIO
from openai import OpenAI
from streamlit_option_menu import option_menu  # Import the option_menu library

# Load environment variables from a .env file for development purposes
load_dotenv()

# Set API keys using environment variables
companies_house_api_key = os.getenv("COMPANIES_HOUSE_API_KEY")
openai_api_key = os.getenv("OPENAI_API_KEY")
news_api_key = os.getenv("NEWS_API_KEY")

# Initialize OpenAI API key
client = OpenAI(api_key=openai_api_key)

# Set page configuration
st.set_page_config(page_title="Company Info App", layout="wide")

if 'current_tab' not in st.session_state:
    st.session_state['current_tab'] = 'Welcome'


def main():
    """
    Main function to run the Streamlit app.
    Initializes session state and handles navigation between tabs.
    """
    # Initialize session state for the current tab if not already set
    if 'current_tab' not in st.session_state:
        st.session_state['current_tab'] = 'Welcome'

    # List of tabs in the app
    tabs = ['Welcome', 'Enter Company Name', 'Results', 'Other Project']
    icons = ['house', 'building', 'bar-chart', 'info-circle']

    # Navigation pane in the sidebar using option_menu
    with st.sidebar:
        st.title("Navigation")
        selected_tab = option_menu(
            menu_title=None,  # Hide the menu title
            options=tabs,
            icons=icons,
            menu_icon="cast",
            default_index=tabs.index(st.session_state['current_tab']),
            styles={
                "container": {"padding": "0!important", "background-color": "#f0f0f0"},
                "icon": {"color": "black", "font-size": "18px"},
                "nav-link": {
                    "font-size": "16px",
                    "text-align": "center",
                    "margin": "5px",
                    "color": "black",
                    "background-color": "#f0f0f0",
                    "border-radius": "5px",
                },
                "nav-link-hover": {"background-color": "#e0e0e0"},
                "nav-link-selected": {"background-color": "#007bff !important", "color": "white !important"},
            },
            key='navigation'
        )
        st.session_state['current_tab'] = selected_tab

    # Display content based on current tab
    if st.session_state['current_tab'] == 'Welcome':
        welcome_tab()
    elif st.session_state['current_tab'] == 'Enter Company Name':
        enter_company_name_tab()
    elif st.session_state['current_tab'] == 'Results':
        results_tab()
    elif st.session_state['current_tab'] == 'Other Project':
        other_project_tab()


def navigate_tabs(direction):
    """
    Function to navigate between tabs using 'Next' and 'Previous' buttons.
    Updates the current tab in session state.
    """
    tabs = ['Welcome', 'Enter Company Name', 'Results', 'Other Project']
    current_index = tabs.index(st.session_state['current_tab'])
    if direction == 'next':
        st.session_state['current_tab'] = tabs[(current_index + 1) % len(tabs)]
    elif direction == 'previous':
        st.session_state['current_tab'] = tabs[(current_index - 1) % len(tabs)]
    st.rerun()


def welcome_tab():
    """
    Displays the content for the 'Welcome' tab.
    Introduces the app and its features.
    """
    st.title("Welcome to the Company Info App")
    st.write("""
    ### Discover Detailed Company Information Easily!

    Welcome to the **Company Info App**! This application allows you to effortlessly search for companies registered in the UK using the Companies House API.

    **What Can You Do With This App?**

    - **Search Companies:** Enter the names of companies you're interested in and select the correct one from a list of suggestions.
    - **View Detailed Information:** Access comprehensive details including company status, incorporation date, registered office address, officers, and more.
    - **Stay Updated:** Read the latest news articles related to the company to stay informed about recent developments.
    - **Generate Summaries:** Get concise summaries of the company's information and recent news.
    - **Download Reports:** Export all the gathered information into a nicely formatted Word document for your records or presentations.

    **Get Started Now!**

    Navigate to the **Enter Company Name** tab to begin your search and explore detailed company profiles.
    """)

    # Navigation buttons
    col1, col2, col3 = st.columns([1, 1, 1])
    with col3:
        if st.button("Next", key="welcome_next"):
            navigate_tabs('next')


def enter_company_name_tab():
    """
    Displays the content for the 'Enter Company Name' tab.
    Allows users to input company names and initiates the selection process.
    """
    st.title("Enter Company Name")

    # Input for company names
    company_input = st.text_area("Enter company names (one per line):", key='company_input')

    if st.button("Search Companies", key="search_companies"):
        company_names = company_input.strip().split('\n')
        company_names = [name.strip() for name in company_names if name.strip()]

        if company_names:
            # Initialize session state variables for the selection process
            st.session_state['company_names'] = company_names
            st.session_state['confirmed_companies'] = {}
            st.session_state['current_company_index'] = 0
            st.session_state['company_selections'] = {}
            st.session_state['selection_complete'] = False
            st.session_state['company_selection_started'] = True
            # Rerun the app to update the interface
            st.rerun()
        else:
            st.warning("Please enter at least one company name.")

    # If the selection process has started, continue it
    if st.session_state.get('company_selection_started', False):
        select_companies()

    # Navigation buttons
    col1, col2, col3 = st.columns([1, 1, 1])
    with col1:
        if st.button("Previous", key="enter_previous"):
            navigate_tabs('previous')
    with col3:
        if st.button("Next", key="enter_next"):
            navigate_tabs('next')


def select_companies():
    """
    Handles the company selection process.
    Displays search results and allows the user to confirm the correct company.
    """
    st.title("Select Companies")

    company_names = st.session_state['company_names']
    total_companies = len(company_names)
    current_index = st.session_state['current_company_index']

    # Progress bar indicating how much of the selection process is complete
    progress = int((current_index / total_companies) * 100)
    st.progress(progress)

    if current_index < total_companies:
        company_name = company_names[current_index]
        st.subheader(f"Results for '{company_name}'")

        # Search Companies House API
        if not companies_house_api_key:
            st.error("Companies House API key is missing.")
            return

        # Prepare authorization for Companies House API
        auth_string = f"{companies_house_api_key}:"
        headers = {
            'Authorization': 'Basic ' + base64.b64encode(auth_string.encode()).decode()
        }
        params = {'q': company_name}

        with st.spinner('Fetching company data...'):
            try:
                response = requests.get(
                    'https://api.company-information.service.gov.uk/search/companies',
                    headers=headers,
                    params=params
                )
                response.raise_for_status()
            except requests.exceptions.RequestException as e:
                st.error(f"Error fetching data for '{company_name}': {e}")
                # Move to next company
                st.session_state['company_selections'][company_name] = None
                st.session_state['current_company_index'] += 1
                st.rerun()
                return

        data = response.json()
        items = data.get('items', [])

        if not items:
            st.write(f"No results found for '{company_name}'.")
            # Move to next company
            st.session_state['company_selections'][company_name] = None
            st.session_state['current_company_index'] += 1
            st.rerun()
            return

        # Show top 5 results as a dropdown
        options = []
        for item in items[:5]:
            options.append(f"{item.get('title')} (Company Number: {item.get('company_number')})")

        # Use a form to group the selectbox and the confirm button
        with st.form(key=f"form_{company_name}"):
            selected = st.selectbox(
                f"Select the correct company for '{company_name}':", options, key=f"select_{company_name}"
            )
            confirm = st.form_submit_button("Confirm Selection")

        if confirm:
            # Store the selected company number
            selected_index = options.index(selected)
            selected_company = items[selected_index]
            company_number = selected_company.get('company_number')

            # Save the selection in session state
            st.session_state['company_selections'][company_name] = company_number
            # Update current company index
            st.session_state['current_company_index'] += 1
            # Check if selection is complete
            if st.session_state['current_company_index'] >= total_companies:
                st.session_state['confirmed_companies'] = st.session_state['company_selections']
                st.session_state['selection_complete'] = True
                st.session_state['company_selection_started'] = False
                st.success("Company selection complete!")
                # Navigate to Results tab
                st.session_state['current_tab'] = 'Results'
            else:
                st.session_state['current_tab'] = 'Enter Company Name'
            # Rerun the app to update
            st.rerun()
    else:
        # All companies have been processed
        st.success("Company selection complete!")
        st.session_state['confirmed_companies'] = st.session_state['company_selections']
        st.session_state['selection_complete'] = True
        st.session_state['company_selection_started'] = False
        # Navigate to Results tab
        st.session_state['current_tab'] = 'Results'
        st.rerun()


def results_tab():
    """
    Displays the results for the selected companies.
    Shows detailed information and provides options to download reports.
    """
    st.title("Results")

    if 'confirmed_companies' not in st.session_state or not st.session_state['confirmed_companies']:
        st.warning("No confirmed companies found. Please go to the 'Enter Company Name' tab to search for companies.")
        return

    companies = st.session_state['confirmed_companies']

    for company_name, company_number in companies.items():
        if company_number is None:
            st.header(f"{company_name} - No Data Available")
            continue
        st.header(f"{company_name} (Company Number: {company_number})")
        display_company_details(company_name, company_number)

    # Navigation buttons
    col1, col2, col3 = st.columns([1, 1, 1])
    with col1:
        if st.button("Previous", key="results_previous"):
            navigate_tabs('previous')
    with col3:
        if st.button("Next", key="results_next"):
            navigate_tabs('next')


def display_company_details(company_name, company_number):
    """
    Fetches and displays detailed information for a single company.
    Includes important information, additional data, recent news, and a summary.
    """
    # Fetch company details from Companies House API
    if not companies_house_api_key:
        st.error("Companies House API key is missing.")
        return

    # Prepare authorization for Companies House API
    auth_string = f"{companies_house_api_key}:"
    headers = {
        'Authorization': 'Basic ' + base64.b64encode(auth_string.encode()).decode()
    }

    with st.spinner('Fetching company details...'):
        try:
            response = requests.get(
                f'https://api.company-information.service.gov.uk/company/{company_number}',
                headers=headers
            )
            response.raise_for_status()
        except requests.exceptions.RequestException as e:
            st.error(f"Error fetching details for '{company_name}': {e}")
            return

    company_data = response.json()

    # Display important information in a styled format
    st.subheader("Important Information")
    important_info = {
        'Company Status': company_data.get('company_status'),
        'Incorporation Date': company_data.get('date_of_creation'),
        'Registered Office Address': company_data.get('registered_office_address')
    }

    for key, value in important_info.items():
        st.write(f"**{key}:** {value}")

    # Additional Information
    st.subheader("Additional Information")
    additional_info_options = ['SIC Codes', 'Accounts', 'Officers']
    selected_info = st.multiselect(
        "Select additional information to view:", additional_info_options, key=f"info_{company_name}"
    )

    if 'SIC Codes' in selected_info:
        sic_codes = company_data.get('sic_codes', [])
        st.write("**SIC Codes:**")
        for code in sic_codes:
            st.write(f"- {code}")
    else:
        sic_codes = []

    if 'Accounts' in selected_info:
        accounts = company_data.get('accounts', {})
        st.write("**Accounts:**")
        st.json(accounts)
    else:
        accounts = {}

    if 'Officers' in selected_info:
        # Fetch officers from Companies House API
        with st.spinner('Fetching officer information...'):
            try:
                officers_response = requests.get(
                    f'https://api.company-information.service.gov.uk/company/{company_number}/officers',
                    headers=headers
                )
                officers_response.raise_for_status()
            except requests.exceptions.RequestException as e:
                st.write("Could not retrieve officers information.")
                officers = []
            else:
                officers_data = officers_response.json()
                officers = officers_data.get('items', [])
                st.write("**Officers:**")
                for officer in officers:
                    st.write(f"- {officer.get('name')} ({officer.get('officer_role')})")
    else:
        officers = []

    # Recent news section using News API
    st.subheader("Recent News")
    news_params = {
        'q': company_name,
        'apiKey': news_api_key,
        'pageSize': 5,
        'sortBy': 'publishedAt'
    }

    with st.spinner('Fetching recent news...'):
        news_response = requests.get('https://newsapi.org/v2/everything', params=news_params)
        if news_response.status_code == 200:
            news_data = news_response.json()
            articles = news_data.get('articles', [])
            if articles:
                for article in articles:
                    st.write(f"- [{article.get('title')}]({article.get('url')})")
            else:
                st.write("No recent news found.")
        else:
            st.write("Error fetching news articles.")
            articles = []
    # Generate 250-word summary using OpenAI API
    st.subheader("Summary")

    summary_key = f"summary_{company_name}"
    if summary_key not in st.session_state:
        with st.spinner('Generating summary...'):
            summary = generate_summary(company_data, articles)
            st.session_state[summary_key] = summary
    else:
        summary = st.session_state[summary_key]
    st.write(summary)

    # Download button for Word document
    if st.button("Download Report as Word Document", key=f"download_{company_name}"):
        generate_word_report(
            company_name=company_name,
            important_info=important_info,
            selected_info=selected_info,
            sic_codes=sic_codes,
            accounts=accounts,
            officers=officers,
            articles=articles
            # Removed the 'summary' parameter
        )


def generate_summary(company_data, articles):
    """
    Generates a summary of the company information and recent news using OpenAI's GPT-4.
    """
    # Prepare the content for the summary
    content = f"Company Information:\n"
    content += f"Status: {company_data.get('company_status')}\n"
    content += f"Incorporation Date: {company_data.get('date_of_creation')}\n"
    content += f"Registered Office Address: {company_data.get('registered_office_address')}\n\n"

    content += "Recent News Articles:\n"
    for article in articles:
        content += f"- {article.get('title')}: {article.get('description')}\n"

    # Use OpenAI's GPT-4 to generate a 250-word summary
    try:
        response = client.chat.completions.create(
            model="gpt-4",
            messages=[
                {"role": "system", "content": "You are a helpful assistant that summarizes company information."},
                {"role": "user", "content": f"Provide a 250-word summary for the following company information and news articles:\n{content}",}
            ],
            max_tokens=400,
            temperature=0.7
        )
        summary = response.choices[0].message.content.strip()

    except Exception as e:
        summary = "Could not generate summary."
        st.error(f"Error generating summary: {e}")

    return summary


def generate_word_report(company_name, important_info, selected_info, sic_codes, accounts, officers, articles):
    """
    Generates a Word document report for the company and provides a download link.
    """
    summary_key = f"summary_{company_name}"
    summary = st.session_state.get(summary_key, "Summary not available.")

    doc = Document()
    doc.add_heading(f"{company_name} Report", 0)

    # Add important information
    doc.add_heading("Important Information", level=1)
    for key, value in important_info.items():
        doc.add_paragraph(f"{key}: {value}")

    # Add additional information
    if selected_info:
        doc.add_heading("Additional Information", level=1)
        if 'SIC Codes' in selected_info and sic_codes:
            doc.add_heading("SIC Codes", level=2)
            for code in sic_codes:
                doc.add_paragraph(code, style='List Bullet')

        if 'Accounts' in selected_info and accounts:
            doc.add_heading("Accounts", level=2)
            doc.add_paragraph(json.dumps(accounts, indent=2))

        if 'Officers' in selected_info and officers:
            doc.add_heading("Officers", level=2)
            for officer in officers:
                doc.add_paragraph(
                    f"{officer.get('name')} ({officer.get('officer_role')})", style='List Bullet'
                )

    # Add recent news
    doc.add_heading("Recent News", level=1)
    if articles:
        for article in articles:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(article.get('title')).bold = True
            p.add_run(f"\n{article.get('description')}\n")
            p.add_run(f"Link: {article.get('url')}")
    else:
        doc.add_paragraph("No recent news articles available.")

    # Add summary
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(summary)

    # Save the document to a BytesIO object
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    # Provide download link
    st.download_button(
        label="Download Word Document",
        data=buffer,
        file_name=f"{company_name}_Report.docx",
        mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )


def other_project_tab():
    """
    Displays the content for the 'Other Project' tab.
    Provides a thank-you message and a link to another project.
    """
    st.title("Other Project")
    st.write("Thank you for trying out my app. I hope you enjoyed it.")
    st.write("Please try out my other project:")
    st.markdown("[Click here to visit my other project](https://app2py-ta4gt3rhhab4to4cqz9hgi.streamlit.app/)")

    # Navigation buttons
    col1, col2, col3 = st.columns([1, 1, 1])
    with col1:
        if st.button("Previous", key="other_previous"):
            navigate_tabs('previous')


if __name__ == "__main__":
    main()
